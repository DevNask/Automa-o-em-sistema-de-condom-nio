import pyautogui as bot
from time import sleep
import pyperclip
from openpyxl import load_workbook

bot.FAILSAFE = True

# =============================
# ABRIR ARQUIVO EXCEL
# =============================

arquivo = "eleicao.xlsx"

workbook = load_workbook(arquivo)
planilha = workbook.active


# -----------------------------
# ABRIR MENU
# -----------------------------
def abrir_menu():

    # Cadastros
    bot.moveTo(x=90, y=281, duration=1)
    bot.click()
    sleep(2)

    # Condomínios
    bot.moveTo(x=83, y=403, duration=1)
    bot.click()
    sleep(2)


# -----------------------------
# CLICAR / PESQUISAR CONDOMÍNIO
# -----------------------------
def clicar_condominio(nome_condominio):

    # Campo de pesquisa
    bot.moveTo(x=495, y=360, duration=1)
    bot.click()

    # Limpa pesquisa anterior
    bot.hotkey('ctrl', 'a')

    # Copia nome do condomínio
    pyperclip.copy(nome_condominio)

    # Cola nome no campo
    bot.hotkey('ctrl', 'v')

    sleep(2)

    # Seleciona o resultado encontrado
    bot.moveTo(x=559, y=409, duration=1)
    bot.click()

    sleep(2)


# -----------------------------
# CAPTURAR TEXTO
# -----------------------------
def capturar_texto():

    sleep(4)

    bot.scroll(+1000)

    sleep(4)

    # -----------------------------------------------
    # Clicar em Corpo Diretivo
    # -----------------------------------------------

    bot.moveTo(x=590, y=349, duration=1)
    bot.click()

    sleep(4)

    # -----------------------------------------------
    # INÍCIO DO MANDATO
    # -----------------------------------------------

    bot.moveTo(x=1103, y=465, duration=1)
    bot.mouseDown()

    bot.moveTo(x=1181, y=466, duration=1)
    bot.mouseUp()

    bot.hotkey('ctrl', 'c')

    sleep(2)

    inicio_mandato = pyperclip.paste()

    print(f"Início: {inicio_mandato}")

    # -----------------------------------------------
    # FIM DO MANDATO
    # -----------------------------------------------

    bot.moveTo(x=1103, y=505, duration=1)
    bot.mouseDown()

    bot.moveTo(x=1178, y=504, duration=1)
    bot.mouseUp()

    bot.hotkey('ctrl', 'c')

    sleep(2)

    fim_mandato = pyperclip.paste()

    print(f"Fim: {fim_mandato}")

    return inicio_mandato, fim_mandato


# -----------------------------
# SALVAR NO EXCEL
# -----------------------------
def salvar_excel(linha, inicio_mandato, fim_mandato):

    # Coluna A = Condomínio
    # Mantém o nome que já estava na planilha

    # Coluna B = Início
    planilha[f"B{linha}"] = inicio_mandato

    # Coluna C = Fim
    planilha[f"C{linha}"] = fim_mandato

    # Salva imediatamente
    workbook.save(arquivo)

    print(f"Salvo na linha {linha}")


# -----------------------------
# VOLTAR PARA LISTA
# -----------------------------
def voltar_lista():

    sleep(3)

    bot.moveTo(x=99, y=241, duration=1)
    bot.click()

    sleep(3)


# -----------------------------
# EXECUÇÃO PRINCIPAL
# -----------------------------

# Começa na linha 2
# A1 = cabeçalho
for linha in range(2, planilha.max_row + 1):

    abrir_menu()

    nome_condominio = planilha[f"A{linha}"].value

    if nome_condominio is None:
        continue

    nome_condominio = str(nome_condominio).strip()

    print(f"Processando linha {linha}: {nome_condominio}")

    clicar_condominio(nome_condominio)

    inicio_mandato, fim_mandato = capturar_texto()

    planilha[f"B{linha}"] = inicio_mandato
    planilha[f"C{linha}"] = fim_mandato

    workbook.save(arquivo)

    voltar_lista()

print("Processo finalizado.")